"""Excel export of any report (see reports.py for the shape). One sheet
per table plus a Summary sheet. Money is written as numbers with a rupee
format so totals still add up in Excel and import cleanly into Tally."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="EEEAFE")
BOLD = Font(bold=True)
MONEY = '"₹"#,##0.00'


def _cell(value, typ):
    if value in (None, ""):
        return None
    if typ == "money":
        return float(Decimal(str(value)))
    if typ in ("int",):
        return int(value)
    if typ == "pct":
        return float(value)
    if typ == "date" and isinstance(value, str) and len(value) == 10:
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _sheet_name(name, used):
    base = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31] or "Sheet"
    n, out = 2, base
    while out in used:
        out = f"{base[:28]} {n}"
        n += 1
    used.add(out)
    return out


def _fit(ws):
    for i, column in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) if c.value is not None else 0) for c in column)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)


def workbook(report, gym_name):
    wb = Workbook()
    used = set()
    ws = wb.active
    ws.title = _sheet_name("Summary", used)
    ws.append([f"{gym_name} · {report['title']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{report['from']} to {report['to']}"])
    ws.append([])
    for s in report["summary"]:
        ws.append([s["label"], _cell(s["value"], s["type"])])
        c = ws.cell(row=ws.max_row, column=2)
        if s["type"] == "money":
            c.number_format = MONEY
        elif s["type"] == "pct":
            c.number_format = '0.0"%"'
    if report.get("notes"):
        ws.append([])
        for n in report["notes"]:
            ws.append([n])
    _fit(ws)
    ws.column_dimensions["A"].width = 34
    for t in report["tables"]:
        sh = wb.create_sheet(_sheet_name(t["name"], used))
        sh.append([c["label"] for c in t["columns"]])
        for c in sh[1]:
            c.font, c.fill, c.alignment = BOLD, HEAD, Alignment(vertical="center")
        for r in t["rows"]:
            sh.append([_cell(r.get(c["key"]), c["type"]) for c in t["columns"]])
        for j, c in enumerate(t["columns"], start=1):
            fmt = {"money": MONEY, "pct": '0.0"%"', "date": "DD-MM-YYYY"}.get(c["type"])
            if fmt:
                for row in sh.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = fmt
        if any(c["type"] == "money" for c in t["columns"]) and t["rows"]:
            last = sh.max_row
            sh.append(["Total"] + [None] * (len(t["columns"]) - 1))
            sh.cell(row=sh.max_row, column=1).font = BOLD
            for j, c in enumerate(t["columns"], start=1):
                if c["type"] == "money" and j > 1:
                    L = get_column_letter(j)
                    cell = sh.cell(row=sh.max_row, column=j, value=f"=SUM({L}2:{L}{last})")
                    cell.number_format, cell.font = MONEY, BOLD
        sh.freeze_panes = "A2"
        _fit(sh)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
