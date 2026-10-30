from datetime import date, timedelta
from io import BytesIO

from django.core.management import call_command

from gym_core.models import CheckIn, GymSettings, Member, Membership
from gym_core.services import today_for
from gym_core.tests import Base

from .models import DailyRun, Reminder, ReminderSettings
from .services import generate


class AutoBase(Base):
    def setUp(self):
        super().setUp()
        self.today = today_for(self.gym)

    def joined(self, name, ends_in, phone=None, paid="1500", **kw):
        """Member on the Monthly plan whose membership ends `ends_in` days from today."""
        m = self.member(name=name, phone=phone, **kw)
        s, b = self.bill(self.rec, m["id"], self.monthly, amount_paid=paid)
        self.assertEqual(s, 201, b)
        Membership.objects.filter(bill_id=b["bill"]["id"]).update(
            start_date=self.today + timedelta(days=ends_in - 29), end_date=self.today + timedelta(days=ends_in),
            original_end_date=self.today + timedelta(days=ends_in))
        return m

    def kinds(self, on=None):
        return sorted((r.kind, r.member.name if r.member_id else "-")
                      for r in Reminder.objects.filter(client=self.gym, for_date=on or self.today))


class DailyListTests(AutoBase):
    def test_renewal_and_expired_reminders(self):
        self.joined("Seven Days", 7)
        self.joined("Ends Today", 0)
        self.joined("Ended Three", -3)
        self.joined("Five Days", 5)
        generate(self.gym, self.today)
        k = self.kinds()
        self.assertIn(("renewal", "Seven Days"), k)
        self.assertIn(("renewal", "Ends Today"), k)
        self.assertIn(("expired", "Ended Three"), k)
        self.assertNotIn(("renewal", "Five Days"), k)
        r = Reminder.objects.get(kind="renewal", member__name="Seven Days")
        self.assertIn("in 7 days", r.text)
        self.assertEqual(str(r.amount), "1500.00")
        self.assertTrue(r.upi_link.startswith("upi://pay?pa=ppfitness%40okaxis"))
        self.assertIn("ppfitness@okaxis", r.text)
        self.assertIn("today", Reminder.objects.get(kind="renewal", member__name="Ends Today").text)

    def test_running_twice_makes_nothing_new_and_renewed_members_are_skipped(self):
        m = self.joined("Seven Days", 7)
        # renewed early: a second membership starting after the first
        s, b = self.bill(self.rec, m["id"], self.monthly, kind="renewal", amount_paid="1500")
        self.assertEqual(s, 201, b)
        self.joined("Other", 7)
        self.assertEqual(generate(self.gym, self.today).get("renewal"), 1)
        self.assertEqual(generate(self.gym, self.today), {})
        self.assertNotIn(("renewal", "Seven Days"), self.kinds())

    def test_dues_reminder_repeats_only_after_the_gap(self):
        self.joined("Owes Money", 20, paid="500")
        generate(self.gym, self.today)
        r = Reminder.objects.get(kind="dues")
        self.assertEqual(str(r.amount), "1000.00")
        self.assertIn("₹ 1,000.00", r.text)
        generate(self.gym, self.today + timedelta(days=3))
        self.assertEqual(Reminder.objects.filter(kind="dues").count(), 1)
        generate(self.gym, self.today + timedelta(days=7))
        self.assertEqual(Reminder.objects.filter(kind="dues").count(), 2)

    def test_birthdays_including_29_feb(self):
        self.member(name="Bday Today", date_of_birth=self.today.replace(year=1995).isoformat()
                    if not (self.today.month == 2 and self.today.day == 29) else "1996-02-29")
        self.member(name="Leap Baby", date_of_birth="1996-02-29")
        generate(self.gym, self.today)
        self.assertIn(("birthday", "Bday Today"), self.kinds())
        generate(self.gym, date(2027, 2, 28))
        self.assertIn(("birthday", "Leap Baby"), self.kinds(date(2027, 2, 28)))

    def test_inactive_call_list(self):
        away = self.joined("Not Coming", 20)
        here = self.joined("Regular", 20)
        CheckIn.objects.create(client=self.gym, member_id=here["id"], local_date=self.today - timedelta(days=1))
        CheckIn.objects.create(client=self.gym, member_id=away["id"], local_date=self.today - timedelta(days=9))
        generate(self.gym, self.today)
        k = self.kinds()
        self.assertIn(("inactive", "Not Coming"), k)
        self.assertNotIn(("inactive", "Regular"), k)
        r = Reminder.objects.get(kind="inactive")
        self.assertEqual(r.days, 9)
        generate(self.gym, self.today + timedelta(days=2))                    # not again within the gap
        self.assertEqual(Reminder.objects.filter(kind="inactive").count(), 1)

    def test_switches_and_settings(self):
        self.joined("Seven Days", 7)
        s, b = self.req("POST", "reminders/settings/", self.owner, {"renewal_on": False, "renewal_days": "10, 2 ,0"})
        self.assertEqual((s, b["settings"]["renewal_days"]), (200, "10,2,0"))
        generate(self.gym, self.today)
        self.assertNotIn("renewal", [k for k, _ in self.kinds()])
        self.assertEqual(self.req("POST", "reminders/settings/", self.owner, {"renewal_days": "7,abc"})[1]["error"], "bad_renewal_days")
        self.assertEqual(self.req("POST", "reminders/settings/", self.rec, {"renewal_on": True})[0], 403)
        # saving the built-in wording unchanged keeps "use default"
        from .texts import RENEWAL
        self.req("POST", "reminders/settings/", self.owner, {"tpl_renewal": RENEWAL})
        self.assertEqual(ReminderSettings.for_client(self.gym).tpl_renewal, "")
        self.req("POST", "reminders/settings/", self.owner, {"tpl_renewal": "Hello {name}, renew {plan} by {end_date}"})
        self.assertEqual(ReminderSettings.for_client(self.gym).tpl_renewal, "Hello {name}, renew {plan} by {end_date}")


class ListApiTests(AutoBase):
    def test_list_builds_itself_and_hides_summary_from_receptionist(self):
        self.joined("Seven Days", 7)
        self.assertFalse(DailyRun.objects.filter(client=self.gym).exists())
        s, b = self.req("GET", "reminders/", self.rec)
        self.assertEqual(s, 200, b)
        self.assertTrue(DailyRun.objects.filter(client=self.gym, date=self.today).exists())
        kinds = {r["kind"] for r in b["reminders"]}
        self.assertIn("renewal", kinds)
        self.assertNotIn("owner_summary", kinds)                               # money summary: owner/manager only
        self.assertIn("owner_summary", {r["kind"] for r in self.req("GET", "reminders/", self.owner)[1]["reminders"]})
        r = b["reminders"][0]
        self.assertTrue(r["whatsapp_link"].startswith("https://wa.me/"))
        self.assertEqual(self.req("GET", "reminders/", self.tr)[0], 403)        # trainers don't send reminders

    def test_mark_sent_skip_call(self):
        self.joined("Seven Days", 7)
        away = self.joined("Away", 20)
        CheckIn.objects.create(client=self.gym, member_id=away["id"], local_date=self.today - timedelta(days=10))
        b = self.req("GET", "reminders/", self.rec)[1]
        ren = [r for r in b["reminders"] if r["kind"] == "renewal"][0]
        s, x = self.req("POST", f"reminders/{ren['id']}/sent/", self.rec, {})
        self.assertEqual((s, x["reminder"]["status"], x["reminder"]["sent_by"]), (200, "sent", "Receptionist"))
        call = [r for r in b["reminders"] if r["kind"] == "inactive"][0]
        self.assertEqual(self.req("POST", f"reminders/{call['id']}/call/", self.rec, {"outcome": "Maybe"})[1]["error"], "bad_outcome")
        s, x = self.req("POST", f"reminders/{call['id']}/call/", self.rec, {"outcome": "Coming back", "notes": "Was travelling"})
        self.assertEqual((x["reminder"]["status"], x["reminder"]["call_outcome"]), ("done", "Coming back"))
        s, x = self.req("POST", f"reminders/{call['id']}/reopen/", self.rec, {})
        self.assertEqual(x["reminder"]["status"], "pending")
        summ = Reminder.objects.get(client=self.gym, kind="owner_summary")
        self.assertEqual(self.req("POST", f"reminders/{summ.id}/sent/", self.rec, {})[0], 404)
        # another gym can't touch this gym's items
        g2 = self.make_gym("Iron House")
        _, o2 = self.user(g2, "owner")
        self.assertEqual(self.req("POST", f"reminders/{ren['id']}/sent/", o2, {})[0], 404)

    def test_owner_summary(self):
        self.joined("Paid Cash", 20)
        s, b = self.req("GET", "summary/", self.owner)
        self.assertEqual(s, 200, b)
        sm = b["summary"]
        self.assertEqual(sm["collected"], "1500.00")
        self.assertEqual(sm["new_members"], 1)
        self.assertIn("Collected: *₹ 1,500.00*", sm["text"])
        self.assertTrue(sm["whatsapp_link"].startswith("https://wa.me/"))
        self.assertEqual(self.req("GET", "summary/", self.rec)[0], 403)

    def test_run_now_and_command(self):
        self.joined("Seven Days", 7)
        s, b = self.req("POST", "reminders/run/", self.rec, {})
        self.assertEqual((s, b["created"].get("renewal")), (200, 1))
        self.assertEqual(self.req("POST", "reminders/run/", self.rec, {"date": (self.today + timedelta(days=1)).isoformat()})[0], 400)
        out = BytesIO()
        from io import StringIO
        buf = StringIO()
        call_command("run_daily_jobs", stdout=buf)
        self.assertIn("Done for", buf.getvalue())


class ReportTests(AutoBase):
    def test_collections_by_mode_and_permissions(self):
        a = self.member(name="A")
        self.bill(self.rec, a["id"], self.monthly, amount_paid="1500", payment_mode="UPI", reference_no="U1")
        b = self.member(name="B")
        self.bill(self.rec, b["id"], self.monthly, amount_paid="1000", payment_mode="Cash")
        s, r = self.req("GET", "reports/collections/", self.owner)
        self.assertEqual(s, 200, r)
        summ = {x["label"]: x["value"] for x in r["report"]["summary"]}
        self.assertEqual((summ["Collected"], summ["UPI"], summ["Cash"]), ("2500.00", "1500.00", "1000.00"))
        self.assertEqual(self.req("GET", "reports/collections/", self.rec)[0], 403)
        self.assertEqual(self.req("GET", "reports/collections/", self.mgr)[0], 200)
        self.assertEqual(self.req("GET", "reports/nope/", self.owner)[0], 404)
        self.assertEqual(self.req("GET", "reports/collections/", self.owner, {"from": "2026-10-10", "to": "2026-10-01"})[0], 400)

    def test_renewals_report(self):
        self.joined("Renewed", 3)
        r1 = Member.objects.get(name="Renewed")
        self.bill(self.rec, str(r1.id), self.monthly, kind="renewal", amount_paid="1500")
        self.joined("Lapsed", 2)
        frm, to = self.today.isoformat(), (self.today + timedelta(days=5)).isoformat()
        r = self.req("GET", "reports/renewals/", self.owner, {"from": frm, "to": to})[1]["report"]
        summ = {x["label"]: x["value"] for x in r["summary"]}
        self.assertEqual((summ["Memberships ending"], summ["Renewed"], summ["Renewal rate"]), (2, 1, 50.0))
        self.assertEqual([x["member"] for x in r["tables"][0]["rows"]], ["Lapsed"])

    def test_gst_report_splits_cgst_sgst(self):
        m = self.member(name="Taxed")
        self.bill(self.rec, m["id"], self.monthly, apply_gst=True, amount_paid="1770")
        n = self.member(name="Untaxed")
        self.bill(self.rec, n["id"], self.monthly, amount_paid="1500")
        r = self.req("GET", "reports/gst/", self.owner)[1]["report"]
        summ = {x["label"]: x["value"] for x in r["summary"]}
        self.assertEqual((summ["Taxable value"], summ["CGST"], summ["SGST"], summ["Bills without GST"]),
                         ("1500.00", "135.00", "135.00", "1500.00"))
        self.assertEqual(len(r["tables"][1]["rows"]), 2)                       # sales register has every bill

    def test_trainer_report_needs_salary_access(self):
        self.assertEqual(self.req("GET", "reports/trainers/", self.mgr)[0], 403)
        self.req("POST", f"team/users/{self.mgr_u.id}/salary-access/", self.owner, {"access": "view"})
        self.assertEqual(self.req("GET", "reports/trainers/", self.mgr)[0], 200)

    def test_profit_and_enquiries(self):
        self.joined("Payer", 20)
        self.req("POST", "expenses/", self.owner, {"date": self.today.isoformat(), "category": "Electricity", "title": "Bill",
                                                   "amount": "400", "mode": "UPI", "status": "Paid"})
        r = self.req("GET", "reports/profit/", self.owner)[1]["report"]
        summ = {x["label"]: x["value"] for x in r["summary"]}
        self.assertEqual((summ["Income"], summ["Expenses paid"], summ["Net"]), ("1500.00", "400.00", "1100.00"))
        self.req("POST", "enquiries/", self.owner, {"name": "Lead", "phone": "9811111111", "source": "Instagram"})
        r = self.req("GET", "reports/enquiries/", self.owner)[1]["report"]
        self.assertEqual(r["tables"][0]["rows"][0]["source"], "Instagram")

    def test_excel_export(self):
        from openpyxl import load_workbook
        m = self.member(name="Taxed")
        self.bill(self.rec, m["id"], self.monthly, apply_gst=True, amount_paid="1770")
        resp = self.client.get("/api/v1/reports/gst/export/", HTTP_AUTHORIZATION=f"Bearer {self.owner}")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Summary", "GST invoices", "Sales register (for Tally)"])
        ws = wb["GST invoices"]
        self.assertEqual(ws["E2"].value, 1500.0)                                # taxable value is a number
        self.assertTrue(str(ws.cell(row=ws.max_row, column=5).value).startswith("=SUM("))
        r = self.client.get("/api/v1/reports/gst/export/", HTTP_AUTHORIZATION=f"Bearer {self.rec}")
        self.assertEqual(r.status_code, 403)

    def test_receipt_pdf(self):
        m = self.member(name="Rohit & Sons <test>")
        s, b = self.bill(self.rec, m["id"], self.monthly, apply_gst=True, amount_paid="1000")
        resp = self.client.get(f"/api/v1/bills/{b['bill']['id']}/receipt.pdf", HTTP_AUTHORIZATION=f"Bearer {self.rec}")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF"))
        r = self.client.get(f"/api/v1/bills/{b['bill']['id']}/receipt.pdf", HTTP_AUTHORIZATION=f"Bearer {self.tr}")
        self.assertEqual(r.status_code, 403)


class LogoAndReceiptTests(AutoBase):
    def png(self, size=(300, 200), fmt="PNG"):
        from PIL import Image
        buf = BytesIO()
        Image.new("RGBA" if fmt == "PNG" else "RGB", size, (108, 92, 231)).save(buf, fmt)
        return buf.getvalue()

    def upload(self, tok, data, name="logo.png", ctype="image/png"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post("/api/v1/gym/settings/logo/", {"logo": SimpleUploadedFile(name, data, ctype)},
                                HTTP_AUTHORIZATION=f"Bearer {tok}")

    def test_logo_upload_resize_and_permissions(self):
        r = self.upload(self.owner, self.png((1600, 800)))
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["settings"]["has_logo"])
        img = self.client.get("/api/v1/gym/logo.png", HTTP_AUTHORIZATION=f"Bearer {self.rec}")
        self.assertEqual((img.status_code, img["Content-Type"]), (200, "image/png"))
        from PIL import Image
        self.assertEqual(Image.open(BytesIO(img.content)).size, (600, 300))          # shrunk, shape kept
        self.assertEqual(self.upload(self.rec, self.png()).status_code, 403)          # owner only
        self.assertEqual(self.upload(self.owner, b"not an image", "x.png").json()["error"], "bad_logo")
        self.assertEqual(self.upload(self.owner, self.png(fmt="GIF"), "x.gif", "image/gif").json()["error"], "bad_logo")
        self.assertEqual(self.upload(self.owner, self.png(fmt="JPEG"), "x.jpg", "image/jpeg").status_code, 200)
        r = self.client.post("/api/v1/gym/settings/logo/remove/", HTTP_AUTHORIZATION=f"Bearer {self.owner}")
        self.assertFalse(r.json()["settings"]["has_logo"])
        self.assertEqual(self.client.get("/api/v1/gym/logo.png", HTTP_AUTHORIZATION=f"Bearer {self.owner}").status_code, 404)
        g2 = self.make_gym("Iron House")                                                # another gym has no logo
        _, o2 = self.user(g2, "owner")
        self.upload(self.owner, self.png())
        self.assertEqual(self.client.get("/api/v1/gym/logo.png", HTTP_AUTHORIZATION=f"Bearer {o2}").status_code, 404)

    def test_html_and_pdf_receipts_show_the_same_things(self):
        self.upload(self.owner, self.png())
        m = self.member(name="Heet Joshi")
        s, b = self.bill(self.owner, m["id"], self.monthly, discount="300", apply_gst=True, amount_paid="1000")
        bid = b["bill"]["id"]
        html = self.client.get(f"/api/v1/bills/{bid}/receipt/", HTTP_AUTHORIZATION=f"Bearer {self.rec}").content.decode()
        self.assertIn('src="data:image/png;base64,', html)
        for text in ("Heet Joshi", "₹ 1,500.00", "− ₹ 300.00", "GST 18% (CGST 9% + SGST 9%)", "₹ 1,416.00",
                     "Balance due", "₹ 416.00", "Partially Paid", "Pay the balance by UPI to"):
            self.assertIn(text, html)
        pdf = self.client.get(f"/api/v1/bills/{bid}/receipt.pdf", HTTP_AUTHORIZATION=f"Bearer {self.rec}").content
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"/Image", pdf)                       # logo drawn
        self.assertIn(b"Inter", pdf)                        # same font as the screen receipt (has the ₹ sign)


class ReceiptLinkTests(AutoBase):
    def test_named_pdf_and_signed_link(self):
        m = self.member(name="Heet Joshi")
        s, b = self.bill(self.rec, m["id"], self.monthly, amount_paid="1000")
        bid = b["bill"]["id"]
        r = self.client.get(f"/api/v1/bills/{bid}/receipt.pdf?download=1", HTTP_AUTHORIZATION=f"Bearer {self.rec}")
        self.assertIn('attachment; filename="RCPT00001 - Heet Joshi.pdf"', r["Content-Disposition"])
        s, j = self.req("GET", f"bills/{bid}/receipt-link/", self.rec)
        self.assertEqual((s, j["filename"]), (200, "RCPT00001 - Heet Joshi.pdf"))
        self.assertTrue(j["url"].endswith("/RCPT00001%20-%20Heet%20Joshi.pdf"))
        pdf = self.client.get(j["url"])                                      # no sign-in needed for the link
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertIn('inline; filename="RCPT00001 - Heet Joshi.pdf"', pdf["Content-Disposition"])
        self.assertIn("attachment", self.client.get(j["download_url"])["Content-Disposition"])
        bad = j["url"].replace("/receipts/", "/receipts/x")
        self.assertEqual(self.client.get(bad).status_code, 404)              # tampered link
        self.assertEqual(self.req("GET", f"bills/{bid}/receipt-link/", self.tr)[0], 403)
        g2 = self.make_gym("Iron House")
        _, o2 = self.user(g2, "owner")
        self.assertEqual(self.req("GET", f"bills/{bid}/receipt-link/", o2)[0], 404)

    def test_link_expires(self):
        from unittest import mock
        m = self.member(name="A / B: C?")
        s, b = self.bill(self.rec, m["id"], self.monthly, amount_paid="1500")
        j = self.req("GET", f"bills/{b['bill']['id']}/receipt-link/", self.rec)[1]
        self.assertEqual(j["filename"], "RCPT00001 - A B C .pdf".replace("C .pdf", "C.pdf"))   # unsafe characters removed
        import time
        real = time.time
        with mock.patch("time.time", lambda: real() + 16 * 60):
            self.assertEqual(self.client.get(j["url"]).status_code, 404)


class DietPdfTests(AutoBase):
    def plan(self, name="Sagar Mhaske"):
        from gym_ops import diet_starter_plans as SP
        m = self.member(name=name)
        key = SP.list_starter_plans("Weight Loss", "Vegetarian")[0]["key"]
        s, b = self.req("POST", "diet/plans/", self.owner, {"member_id": m["id"], "starter_key": key})
        self.assertEqual(s, 201, b)
        return b["plan"]

    def test_pdf_and_staff_link(self):
        p = self.plan()
        r = self.client.get(f"/api/v1/diet/plans/{p['id']}/pdf/?download=1", HTTP_AUTHORIZATION=f"Bearer {self.owner}")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn('attachment; filename="Diet plan - Sagar Mhaske - v1.pdf"', r["Content-Disposition"])
        s, j = self.req("GET", f"diet/plans/{p['id']}/pdf-link/", self.rec)
        self.assertEqual(s, 403)                                  # receptionists have no diet access
        s, j = self.req("GET", f"diet/plans/{p['id']}/pdf-link/", self.owner)
        self.assertEqual(self.client.get(j["url"]).status_code, 200)
        from unittest import mock
        import time
        real = time.time
        with mock.patch("time.time", lambda: real() + 20 * 60):  # staff links last 15 minutes
            self.assertEqual(self.client.get(j["url"]).status_code, 404)

    def test_whatsapp_message_has_member_pdf_link_for_that_version(self):
        p = self.plan()
        s, j = self.req("POST", f"diet/plans/{p['id']}/share/", self.owner, {})
        self.assertEqual(s, 200, j)
        self.assertIn("Full plan with calories (PDF):", j["whatsapp"]["text"])
        self.assertIn(j["pdf"]["url"], j["whatsapp"]["text"])
        from .views import is_local
        self.assertTrue(is_local("http://127.0.0.1:8000") and is_local("http://localhost:8000"))
        self.assertFalse(is_local("https://fitdesk.onrender.com"))
        url = j["pdf"]["url"].split("testserver", 1)[1]
        from unittest import mock
        import time
        real = time.time
        with mock.patch("time.time", lambda: real() + 30 * 86400):  # member link still works after a month
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200)
            self.assertIn("Diet plan - Sagar Mhaske - v1.pdf", r["Content-Disposition"])
        # revising the plan doesn't change what the member was sent
        full = self.req("GET", f"diet/plans/{p['id']}/", self.owner)[1]["plan"]
        s, rev = self.req("POST", f"diet/plans/{p['id']}/", self.owner, {"name": "Changed", "meals": full["version"]["meals"][:1],
                                                                          "change_note": "fewer meals"})
        self.assertEqual(rev["plan"]["current_version"], 2)
        self.assertEqual(self.client.get(url).status_code, 200)
        with mock.patch("time.time", lambda: real() + 61 * 86400):
            self.assertEqual(self.client.get(url).status_code, 404)
        with mock.patch.dict("os.environ", {"PUBLIC_BASE_URL": "https://fitdesk.example.com"}):
            j = self.req("POST", f"diet/plans/{p['id']}/share/", self.owner, {})[1]
        self.assertTrue(j["pdf"]["url"].startswith("https://fitdesk.example.com/api/v1/d/"))
        self.assertFalse(j["pdf"]["local_only"])
        self.assertNotIn("PDF", self.req("POST", f"diet/plans/{p['id']}/share/", self.owner, {"include_pdf": False})[1]["whatsapp"]["text"])

    def test_trainer_only_own_clients(self):
        p = self.plan()
        self.assertEqual(self.req("GET", f"diet/plans/{p['id']}/pdf-link/", self.tr)[0], 404)
        g2 = self.make_gym("Iron House")
        _, o2 = self.user(g2, "owner")
        self.assertEqual(self.client.get(f"/api/v1/diet/plans/{p['id']}/pdf/", HTTP_AUTHORIZATION=f"Bearer {o2}").status_code, 404)


class GstInclusivePlanTests(AutoBase):
    def test_inclusive_split(self):
        from decimal import Decimal as D
        from gym_core import rules as R
        t = R.bill_totals_inclusive("1500", "0", "18")
        self.assertEqual((t["net"], t["gst_amount"], t["total"]), (D("1271.19"), D("228.81"), D("1500.00")))
        t = R.bill_totals_inclusive("1180", "180", "18")          # discount comes off the inclusive price
        self.assertEqual((t["net"], t["gst_amount"], t["total"]), (D("847.46"), D("152.54"), D("1000.00")))

    def test_plan_with_gst_included_bills_without_adding_gst(self):
        s, p = self.req("POST", "plans/", self.owner, {"name": "Monthly incl GST", "duration_days": 30, "fee": "1500",
                                                       "gst_inclusive": True})
        self.assertEqual(s, 201, p)
        self.assertTrue(p["plan"]["gst_inclusive"])
        s, lst = self.req("GET", "plans/", self.rec)
        self.assertEqual(lst["gst_percent"], "18.00")
        m = self.member()
        s, q = self.req("POST", "bills/quote/", self.rec, {"member_id": m["id"], "plan_id": p["plan"]["id"], "apply_gst": False,
                                                           "amount_paid": "1500"})
        q = q["quote"]
        self.assertEqual((q["total"], q["gst_amount"], q["taxable"], q["gst_inclusive"], q["payment_status"]),
                         ("1500.00", "228.81", "1271.19", True, "Paid"))
        s, b = self.bill(self.rec, m["id"], p["plan"]["id"], amount_paid="1500")
        self.assertEqual(s, 201, b)
        self.assertEqual((b["bill"]["total"], b["bill"]["gst_amount"], b["bill"]["gst_inclusive"]), ("1500.00", "228.81", True))
        s, html = self.req("GET", f"bills/{b['bill']['id']}/receipt/", self.rec)
        self.assertIn("GST 18% included", html)
        self.assertIn("Price before GST", html)
        r = self.client.get(f"/api/v1/bills/{b['bill']['id']}/receipt.pdf", HTTP_AUTHORIZATION=f"Bearer {self.rec}")
        self.assertEqual(r.status_code, 200)
        # GST report counts the GST inside the price
        s, rep = self.req("GET", "reports/gst/", self.owner)
        self.assertEqual(s, 200, rep)

    def test_gst_inclusive_needs_gst_percent(self):
        GymSettings.objects.filter(client=self.gym).update(gst_percent=0)
        s, p = self.req("POST", "plans/", self.owner, {"name": "X", "duration_days": 30, "fee": "1500", "gst_inclusive": True})
        self.assertEqual((s, p["error"]), (400, "gst_not_set_up"))

    def test_normal_plan_unchanged(self):
        m = self.member()
        s, q = self.req("POST", "bills/quote/", self.rec, {"member_id": m["id"], "plan_id": self.monthly, "apply_gst": True,
                                                           "amount_paid": "0"})
        self.assertEqual((q["quote"]["total"], q["quote"]["gst_amount"], q["quote"]["gst_inclusive"]), ("1770.00", "270.00", False))


class DietChangeRemoveTests(DietPdfTests):
    def give(self, m, replace=None, **kw):
        from gym_ops import diet_starter_plans as SP
        key = SP.list_starter_plans("Muscle Gain", "Vegetarian")[0]["key"]
        b = {"member_id": m["id"], "starter_key": key, **kw}
        if replace is not None:
            b["replace_active"] = replace
        s, j = self.req("POST", "diet/plans/", self.owner, b)
        self.assertEqual(s, 201, j)
        return j

    def status(self, pid):
        return self.req("GET", f"diet/plans/{pid}/", self.owner)[1]["plan"]["status"]

    def test_replace_keep_both_remove_restore(self):
        old = self.plan()
        m = old["member"]
        both = self.give(m)                                   # keep both
        self.assertEqual((both["replaced"], self.status(old["id"])), (0, "Active"))
        new = self.give(m, replace=True)                      # replace: other active plans become Completed
        self.assertEqual(new["replaced"], 2)
        self.assertEqual((self.status(old["id"]), self.status(both["plan"]["id"]), self.status(new["plan"]["id"])),
                         ("Completed", "Completed", "Active"))
        # remove = archive; hidden from the default list, shown under Archived, can be restored
        s, j = self.req("POST", f"diet/plans/{new['plan']['id']}/", self.owner, {"status": "Archived"})
        self.assertEqual(s, 200, j)
        ids = [p["id"] for p in self.req("GET", "diet/plans/", self.owner)[1]["plans"]]
        self.assertNotIn(new["plan"]["id"], ids)
        ids = [p["id"] for p in self.req("GET", "diet/plans/?view=Archived", self.owner)[1]["plans"]]
        self.assertEqual(ids, [new["plan"]["id"]])
        s, j = self.req("POST", f"diet/plans/{new['plan']['id']}/", self.owner, {"status": "Active"})
        self.assertEqual(j["plan"]["status"], "Active")

    def test_draft_made_active_can_replace(self):
        old = self.plan()
        d = self.give(old["member"], status="Draft")
        s, j = self.req("POST", f"diet/plans/{d['plan']['id']}/", self.owner, {"status": "Active", "replace_active": True})
        self.assertEqual((s, j["plan"]["status"], self.status(old["id"])), (200, "Active", "Completed"))


class MemberDuplicateTests(AutoBase):
    def test_phone_with_trunk_zero_is_same_number(self):
        from gym_core.services import clean_phone
        self.assertEqual(clean_phone("09028644659"), "+919028644659")
        self.assertEqual(clean_phone("+09028644659"), "+919028644659")
        self.assertEqual(clean_phone("00919028644659"), "+919028644659")
        self.member(name="Sanket Shrikant", phone="9028644659")
        s, j = self.req("POST", "members/", self.rec, {"name": "Sanket Shrikant", "phone": "09028644659", "gender": "Male"})
        self.assertEqual((s, j["error"]), (409, "phone_taken"))
        self.assertIn("Sanket Shrikant", j["message"])

    def test_delete_only_when_nothing_linked(self):
        m = self.member(name="Typo Entry")
        s, j = self.req("POST", f"members/{m['id']}/delete/", self.rec, {})
        self.assertEqual(s, 403)                                  # receptionists can't remove members
        s, j = self.req("GET", f"members/{m['id']}/records/", self.owner)
        self.assertEqual(j["records"], {})
        s, j = self.req("POST", f"members/{m['id']}/delete/", self.owner, {})
        self.assertEqual(s, 200, j)
        self.assertFalse(Member.objects.filter(pk=m["id"]).exists())
        paid = self.joined("Has Bill", 10)
        s, j = self.req("POST", f"members/{paid['id']}/delete/", self.owner, {})
        self.assertEqual((s, j["error"], j["records"]["bills"]), (409, "member_has_records", 1))

    def test_merge_moves_history_and_removes_duplicate(self):
        keep = self.member(name="Sanket Shrikant", phone="9028644659")
        Member.objects.filter(pk=keep["id"]).update(email="")
        dup = self.joined("Sanket S", 20, phone="9028644600")
        Member.objects.filter(pk=dup["id"]).update(email="sanket@example.com")
        s, j = self.req("POST", f"members/{dup['id']}/merge/", self.owner, {"into_id": keep["id"]})
        self.assertEqual(s, 200, j)
        self.assertEqual((j["moved"]["bills"], j["moved"]["memberships"]), (1, 1))
        self.assertEqual(j["member"]["email"], "sanket@example.com")
        self.assertEqual(j["member"]["status"], "active")
        self.assertFalse(Member.objects.filter(pk=dup["id"]).exists())
        s, j = self.req("POST", f"members/{keep['id']}/merge/", self.owner, {"into_id": keep["id"]})
        self.assertEqual(s, 400)


class EmailTests(AutoBase):
    def setUp(self):
        super().setUp()
        from django.core import mail
        self.mail = mail

    def setup_email(self):
        s, j = self.req("POST", "email/settings/", self.owner, {"mode": "own", "provider": "gmail",
                                                                 "sender_email": "HMFitness@gmail.com", "password": "abcd efgh ijkl mnop"})
        self.assertEqual(s, 200, j)
        return j["settings"]

    def test_settings_hide_password_and_need_owner(self):
        s, j = self.req("POST", "email/settings/", self.rec, {"sender_email": "x@y.com"})
        self.assertEqual(s, 403)
        st = self.setup_email()
        self.assertTrue(st["has_password"] and st["ready"])
        self.assertEqual((st["smtp_host"], st["smtp_port"], st["sender_email"]), ("smtp.gmail.com", 587, "hmfitness@gmail.com"))
        self.assertNotIn("password", st)
        from .mailer import decrypt
        from .models import EmailSettings
        es = EmailSettings.objects.get(client=self.gym)
        self.assertNotIn("abcd", es.password_enc)
        self.assertEqual(decrypt(es.password_enc), "abcdefghijklmnop")        # Gmail app password spaces removed

    def test_not_set_up(self):
        m = self.joined("Heet Joshi", 10)
        bill = self.req("GET", f"bills/?member_id={m['id']}", self.rec)[1]["bills"][0]
        s, j = self.req("POST", f"bills/{bill['id']}/email/", self.rec, {"to": "heet@example.com"})
        self.assertEqual((s, j["error"]), (400, "email_not_set_up"))

    def test_receipt_diet_reminder_and_test_email(self):
        self.setup_email()
        s, j = self.req("POST", "email/test/", self.owner, {})
        self.assertEqual(s, 200, j)
        self.assertTrue(j["settings"]["verified_at"])
        m = self.joined("Heet Joshi", 3)
        bill = self.req("GET", f"bills/?member_id={m['id']}", self.rec)[1]["bills"][0]
        s, j = self.req("POST", f"bills/{bill['id']}/email/", self.rec, {})
        self.assertEqual((s, j["error"]), (400, "no_email"))
        s, j = self.req("POST", f"bills/{bill['id']}/email/", self.rec, {"to": "heet@example.com", "save_to_member": True})
        self.assertEqual(s, 200, j)
        msg = self.mail.outbox[-1]
        self.assertEqual(msg.to, ["heet@example.com"])
        self.assertIn(bill["receipt_no"], msg.subject)
        self.assertEqual(msg.attachments[0][0], f"{bill['receipt_no']} - Heet Joshi.pdf")
        self.assertTrue(msg.attachments[0][1].startswith(b"%PDF"))
        self.assertIn("hmfitness@gmail.com", msg.from_email)
        self.assertEqual(Member.objects.get(pk=m["id"]).email, "heet@example.com")
        # next time the saved address is used
        s, j = self.req("POST", f"bills/{bill['id']}/email/", self.rec, {})
        self.assertEqual(j["to"], "heet@example.com")
        # diet plan
        from gym_ops import diet_starter_plans as SP
        key = SP.list_starter_plans("Weight Loss", "Vegetarian")[0]["key"]
        p = self.req("POST", "diet/plans/", self.owner, {"member_id": m["id"], "starter_key": key})[1]["plan"]
        s, j = self.req("POST", f"diet/plans/{p['id']}/email/", self.owner, {})
        self.assertEqual(s, 200, j)
        self.assertTrue(self.mail.outbox[-1].attachments[0][0].startswith("Diet plan - Heet Joshi"))
        # reminder
        generate(self.gym, self.today)
        r = Reminder.objects.filter(client=self.gym, member_id=m["id"], kind="renewal").first()
        s, j = self.req("POST", f"reminders/{r.id}/email/", self.rec, {})
        self.assertEqual(s, 200, j)
        self.assertEqual((j["reminder"]["status"], j["reminder"]["channel"]), ("sent", "email"))
        s, j = self.req("GET", "email/log/", self.owner)
        self.assertEqual([e["kind"] for e in j["emails"]][:4], ["reminder", "diet_plan", "receipt", "receipt"])

    def test_bad_password_gives_plain_reason(self):
        import smtplib
        from unittest import mock
        self.setup_email()
        m = self.joined("Heet Joshi", 10)
        bill = self.req("GET", f"bills/?member_id={m['id']}", self.rec)[1]["bills"][0]
        with mock.patch("django.core.mail.EmailMultiAlternatives.send", side_effect=smtplib.SMTPAuthenticationError(535, b"bad")):
            s, j = self.req("POST", f"bills/{bill['id']}/email/", self.rec, {"to": "heet@example.com"})
        self.assertEqual((s, j["error"]), (400, "email_failed"))
        self.assertIn("app password", j["message"])
        from .models import EmailLog
        self.assertEqual(EmailLog.objects.get(client=self.gym).status, "failed")
